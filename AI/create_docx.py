import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, hex_color="E2E8F0"):
    """Sets light subtle borders on table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def add_inline_formatted_text(paragraph, text):
    """Parses markdown inline bold, code, and italic."""
    # Pattern to match bold-italic, bold, inline code, and italic
    # Tokens: ***...***, **...**, `...`, *...*
    token_pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|`.*?`|\*.*?\*)')
    tokens = token_pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 25, 60)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            # Check for links [text](url)
            sub_tokens = re.split(r'(\[.*?\]\(.*?\))', token)
            for st in sub_tokens:
                link_match = re.match(r'\[(.*?)\]\((.*?)\)', st)
                if link_match:
                    link_text = link_match.group(1)
                    run = paragraph.add_run(link_text)
                    run.font.color.rgb = RGBColor(37, 99, 235)
                    run.underline = True
                else:
                    paragraph.add_run(st)

def render_code_block(doc, code_lines, lang=""):
    """Renders a code block inside a shaded, bordered single-cell table."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left accent border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    # Spacing after table
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(4)

def render_table(doc, table_rows):
    """Renders a parsed markdown table."""
    if not table_rows:
        return
    
    # Filter out empty or delimiter rows
    cleaned_rows = []
    for r in table_rows:
        cells = [c.strip() for c in r.strip().strip('|').split('|')]
        # Check if delimiter row like :--- or ---
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        cleaned_rows.append(cells)
    
    if not cleaned_rows:
        return
    
    col_count = max(len(r) for r in cleaned_rows)
    # Normalize rows
    for r in cleaned_rows:
        while len(r) < col_count:
            r.append("")
            
    doc_table = doc.add_table(rows=len(cleaned_rows), cols=col_count)
    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc_table.autofit = True
    
    for row_idx, row_data in enumerate(cleaned_rows):
        is_header = (row_idx == 0)
        for col_idx, cell_value in enumerate(row_data):
            cell = doc_table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            
            if is_header:
                set_cell_background(cell, "1E3A8A") # Navy header
                run = p.add_run(cell_value)
                run.bold = True
                run.font.name = 'Segoe UI'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg_color = "F8FAFC" if (row_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_borders(cell, "E2E8F0")
                add_inline_formatted_text(p, cell_value)
                for r in p.runs:
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(9.0)
                    if not r.font.color.rgb:
                        r.font.color.rgb = RGBColor(51, 65, 85)

    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(6)

def convert_md_file_to_docx(doc, filepath, is_first_file=False):
    """Parses a markdown file and appends formatted elements to doc."""
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    code_lang = ""
    
    in_table = False
    table_lines = []
    
    for line in lines:
        line_str = line.rstrip('\r\n')
        stripped = line_str.strip()
        
        # Handle Code Block boundary
        if stripped.startswith('```'):
            if in_code_block:
                render_code_block(doc, code_lines, code_lang)
                in_code_block = False
                code_lines = []
                code_lang = ""
            else:
                if in_table:
                    render_table(doc, table_lines)
                    in_table = False
                    table_lines = []
                in_code_block = True
                code_lang = stripped[3:].strip()
            continue
            
        if in_code_block:
            code_lines.append(line_str)
            continue
            
        # Handle Table rows
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                render_table(doc, table_lines)
                in_table = False
                table_lines = []
                
        # Empty lines
        if not stripped:
            continue
            
        # Horizontal rules
        if stripped in ('---', '***', '___'):
            # subtle spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            continue
            
        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[2:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Navy
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[3:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235) # Blue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[4:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59) # Slate
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[5:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            
        # Bullet list
        elif re.match(r'^[\*\-\+]\s+', stripped):
            item_text = re.sub(r'^[\*\-\+]\s+', '', stripped)
            # Checkbox items
            if item_text.startswith('[ ] '):
                item_text = '☐ ' + item_text[4:]
            elif item_text.startswith('[x] ') or item_text.startswith('[X] '):
                item_text = '☑ ' + item_text[4:]
                
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, item_text)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.0)
                if not r.font.color.rgb:
                    r.font.color.rgb = RGBColor(51, 65, 85)
                    
        # Numbered list
        elif re.match(r'^\d+\.\s+', stripped):
            item_text = stripped
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, item_text)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.0)
                if not r.font.color.rgb:
                    r.font.color.rgb = RGBColor(51, 65, 85)
                    
        # Blockquote
        elif stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("┃ ")
            run.font.color.rgb = RGBColor(37, 99, 235)
            run.bold = True
            add_inline_formatted_text(p, quote_text)
            for r in p.runs[1:]:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(9.5)
                r.italic = True
                r.font.color.rgb = RGBColor(71, 85, 105)
                
        # Normal Body Paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, stripped)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.0)
                if not r.font.color.rgb:
                    r.font.color.rgb = RGBColor(51, 65, 85)
                    
    # Clean up trailing elements if any
    if in_code_block and code_lines:
        render_code_block(doc, code_lines, code_lang)
    if in_table and table_lines:
        render_table(doc, table_lines)

def build_master_document(notes_dir, output_docx_path):
    print("Initializing Word Document...")
    doc = Document()
    
    # Page Setup: Standard Letter, 0.8 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header & Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("GenAI Comprehensive Study Notes & Interview Guide  |  Page ")
        f_run.font.name = 'Segoe UI'
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(148, 163, 184)
        
        # Add Page Number field XML
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        f_p._p.append(fldSimple)

    # 1. COVER PAGE
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(72)
    title_p.paragraph_format.space_after = Pt(12)
    t_run = title_p.add_run("GENAI & MACHINE LEARNING\nMASTER STUDY COMPENDIUM")
    t_run.font.name = 'Segoe UI'
    t_run.font.size = Pt(28)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(30, 58, 138)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(6)
    sub_p.paragraph_format.space_after = Pt(24)
    s_run = sub_p.add_run("Complete Technical Reference Guide from ML Foundations to Frontier LLMs, Advanced RAG, Autonomous Agents & Production Serving")
    s_run.font.name = 'Segoe UI'
    s_run.font.size = Pt(13)
    s_run.font.color.rgb = RGBColor(71, 85, 105)
    
    meta_table = doc.add_table(rows=4, cols=2)
    meta_data = [
        ("Curriculum Scope", "1-Month Intensive GenAI Syllabus (28 Days)"),
        ("Coverage", "Theory, Math Derivations, PyTorch Code, Architecture & Top 50 Interview Questions"),
        ("Key Technologies", "PyTorch, Transformers, LoRA/QLoRA, LangGraph, LlamaIndex, Neo4j, vLLM"),
        ("Target Use Cases", "1-Day Rapid Revision, Periodic Concept Refresh, Technical Interview Prep")
    ]
    for idx, (k, v) in enumerate(meta_data):
        cell_k = meta_table.cell(idx, 0)
        cell_v = meta_table.cell(idx, 1)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, 60, 60, 100, 100)
        set_cell_margins(cell_v, 60, 60, 100, 100)
        
        pk = cell_k.paragraphs[0]
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.name = 'Segoe UI'
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = RGBColor(30, 41, 59)
        
        pv = cell_v.paragraphs[0]
        rv = pv.add_run(v)
        rv.font.name = 'Segoe UI'
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = RGBColor(51, 65, 85)
        
    doc.add_page_break()
    
    # 2. Iterate through all markdown files in chronological order
    md_files = sorted([f for f in os.listdir(notes_dir) if f.endswith('.md')])
    print(f"Discovered {len(md_files)} markdown files to process...")
    
    for i, filename in enumerate(md_files):
        filepath = os.path.join(notes_dir, filename)
        print(f"[{i+1}/{len(md_files)}] Processing: {filename} ...")
        
        convert_md_file_to_docx(doc, filepath)
        
        # Add page break between files (except last)
        if i < len(md_files) - 1:
            doc.add_page_break()
            
    print(f"Saving combined master document to: {output_docx_path} ...")
    doc.save(output_docx_path)
    print("SUCCESS: Master DOCX document generated successfully!")

if __name__ == "__main__":
    NOTES_DIR = r"c:\Users\arman\Downloads\SM\GenAI_Study_Notes"
    OUTPUT_DOCX = r"c:\Users\arman\Downloads\SM\GenAI_Complete_Study_Notes.docx"
    build_master_document(NOTES_DIR, OUTPUT_DOCX)
